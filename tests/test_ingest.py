from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

import ingest
import pytest
from openpyxl import Workbook

EVALS = Path(__file__).resolve().parents[1] / "evals"


def run(folder: Path, out: Path) -> tuple[list[dict[str, Any]], str]:
    assert ingest.main([str(folder), "--out", str(out)]) == 0
    facts = [json.loads(line) for line in (out / "facts.jsonl").read_text().splitlines()]
    return facts, (out / "conflicts.md").read_text()


@pytest.mark.parametrize(
    "case,needle",
    [
        ("grocer-costout", "revenue"),
        ("market-entry", "market"),
        ("steerco-update", "go-live"),
        ("take-private-valuation", "diluted, shares"),
        ("sell-side-pitch", "adjusted, ebitda"),
    ],
)
def test_planted_conflict_found_in_each_fixture(tmp_path: Path, case: str, needle: str) -> None:
    facts, conflicts = run(EVALS / "cases" / case / "input", tmp_path / "brief")
    lines = [ln for ln in conflicts.splitlines() if ln.startswith("- **")]
    assert len(lines) == 1, conflicts
    assert needle in lines[0]
    assert all(f["locator"] and f["file"] for f in facts)


def test_independent_conflict_and_non_conflict(tmp_path: Path) -> None:
    src = tmp_path / "src"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L ($M)"
    ws.append(["Metric", "FY24"])
    ws.append(["Revenue", 1250])
    ws.append(["Headcount", 3400])
    wb.save(src / "model.xlsx")
    (src / "memo.md").write_text("Revenue reached $1.25B in FY24. Headcount is 3,900 people.\n")
    facts, conflicts = run(src, tmp_path / "brief")
    assert "revenue" not in conflicts  # $1.25B equals 1,250 $M: same figure at a different scale
    assert "employees" in conflicts  # 3,400 vs 3,900 is a real conflict (headcount -> employees)


def test_formula_evaluator_is_arithmetic_only() -> None:
    grid = {"A1": 2.0, "A2": 3.0, "A3": 5.0}
    look = lambda ref: [grid[r] for r in ([ref] if ":" not in ref else ["A1", "A2", "A3"])]  # noqa: E731
    assert ingest.Formula("=SUM(A1:A3)*2-A1/2", look).value() == 19.0
    assert ingest.Formula("=(A1+A2)*A3", look).value() == 25.0
    with pytest.raises(ValueError):
        ingest.Formula("=__import__('os')", look).value()
    with pytest.raises(ValueError):
        ingest.Formula("=VLOOKUP(A1,A2:A3,1)", look).value()


def test_circular_reference_does_not_hang(tmp_path: Path) -> None:
    src = tmp_path / "src"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "x", "=B2+1", "y", "=B1+1"
    wb.save(src / "loop.xlsx")
    facts, _ = run(src, tmp_path / "brief")
    assert not [f for f in facts if f["locator"] in ("Sheet!B1", "Sheet!B2")]


def test_symlinks_hidden_and_oversize_files_are_not_read(tmp_path: Path) -> None:
    src = tmp_path / "src"
    src.mkdir()
    secret = tmp_path / "secret.txt"
    secret.write_text("Password budget $999M\n")
    os.symlink(secret, src / "link.txt")
    (src / ".hidden.md").write_text("Revenue $5M\n")
    (src / "big.txt").write_text("x" * 2048)
    assert ingest.main([str(src), "--out", str(tmp_path / "b"), "--max-mb", "0.001"]) == 0
    sources = (tmp_path / "b" / "sources.md").read_text()
    facts = (tmp_path / "b" / "facts.jsonl").read_text()
    assert "999" not in facts and "symlink" in sources
    assert ".hidden" not in sources
    assert "size cap" in sources


def test_output_inside_input_is_refused(tmp_path: Path) -> None:
    src = tmp_path / "src"
    src.mkdir()
    assert ingest.main([str(src), "--out", str(src / "brief")]) == 2


def test_symlinked_input_folder_is_refused(tmp_path: Path) -> None:
    real = tmp_path / "real"
    real.mkdir()
    os.symlink(real, tmp_path / "alias")
    assert ingest.main([str(tmp_path / "alias"), "--out", str(tmp_path / "b")]) == 2


def test_labels_like_q1_and_fy25_are_not_figures() -> None:
    vals = [v for v, *_ in ingest.numbers_in("Q1 FY25 revenue rose 12% to $4.2M in H2")]
    assert vals == [12.0, 4.2]


def test_hyphen_between_numbers_is_a_range_not_a_minus() -> None:
    vals = [v for v, *_ in ingest.numbers_in("overstaffed by 15-20% on Tuesdays; margin fell -3%")]
    assert vals == [15.0, 20.0, -3.0]


def test_percent_ranges_and_year_like_values_in_tables(tmp_path: Path) -> None:
    assert [v for v, *_ in ingest.numbers_in("WACC 8.5%-9.5%; growth 2.0%-3.0%")] == [8.5, 9.5, 2.0, 3.0]
    src = tmp_path / "src"
    src.mkdir()
    (src / "comps.csv").write_text("company,ev_usd_m,year\nDogwood,1980,2024\n")
    facts, _ = run(src, tmp_path / "brief")
    evs = [f for f in facts if "ev_usd_m" in f["text"]]
    assert evs and evs[0]["value"] == 1980.0
    assert not [f for f in facts if "| year" in f["text"]]


def test_integer_followed_by_a_word_is_kept() -> None:
    found = ingest.numbers_in("lift margin by 2 points; only 3 of 10 hotels; 142 stores")
    assert [v for v, *_ in found] == [2, 3, 10, 142]
    assert [(v, u) for v, u, *_ in ingest.numbers_in("revenue $4.8 million and 12 x-rays")] == [(4.8, "USD"), (12, "")]


def test_zip_bomb_part_and_corrupt_pdf_land_in_not_read(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import zipfile

    src = tmp_path / "src"
    src.mkdir()
    with zipfile.ZipFile(src / "bomb.docx", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w:document>" + "A" * 200_000 + "</w:document>")
    (src / "broken.pdf").write_bytes(b"%PDF-1.4\n1 0 obj << /Type /Catalog >> garbage\n%%EOF")
    (src / "notes.md").write_text("Revenue was $12.5M in FY25.\n")
    monkeypatch.setattr(ingest, "MAX_MEMBER_BYTES", 100_000)  # the 200 KB part stands in for a bomb
    facts, _ = run(src, tmp_path / "brief")
    sources = (tmp_path / "brief" / "sources.md").read_text()
    assert "bomb.docx" in sources and "over the" in sources  # refused, and reported
    assert "broken.pdf" in sources  # reported, not fatal
    assert any(f["file"] == "notes.md" for f in facts)  # the rest of the folder was still read
