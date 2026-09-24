from __future__ import annotations

import importlib.util
import sys
import types
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "premium-decks" / "scripts"
sys.path.insert(0, str(SCRIPTS))

# make_pptx(slides, name="deck.pptx") -> path of the saved .pptx
MakePptx = Callable[..., Path]


def load_script(name: str) -> types.ModuleType:
    """Import a hyphenated script (e.g. storyline-lint.py) as a module."""
    path = SCRIPTS / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name.replace("-", "_"), path)
    assert spec is not None and spec.loader is not None, f"cannot load {path}"
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture
def make_pptx(tmp_path: Path) -> MakePptx:
    """Build a small .pptx: slides = [{title, body, notes, chart, table, boxes}]."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    # python-pptx ships CategoryChartData unannotated; state the signature used here
    new_chart_data: Callable[[], Any] = CategoryChartData

    def build(slides: list[dict[str, Any]], name: str = "deck.pptx") -> Path:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        blank = prs.slide_layouts[6]
        for spec in slides:
            s = prs.slides.add_slide(blank)
            ty = spec.get("title_y", 0.5)
            tb = s.shapes.add_textbox(Inches(0.6), Inches(ty), Inches(12), Inches(1.0))
            tb.name = "Title 1"
            tb.text_frame.text = spec.get("title", "")
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28) if spec.get("title") else None
            if spec.get("body"):
                b = s.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(6), Inches(3))
                b.name = "Body"
                b.text_frame.text = spec["body"]
                b.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
            for box in spec.get("boxes", []):
                t = s.shapes.add_textbox(Inches(box["x"]), Inches(box["y"]), Inches(box["w"]), Inches(box["h"]))
                t.name = box.get("name", "Box")
                t.text_frame.word_wrap = True
                t.text_frame.auto_size = MSO_AUTO_SIZE.NONE
                t.text_frame.text = box["text"]
                t.text_frame.paragraphs[0].runs[0].font.size = Pt(box.get("pt", 14))
                if box.get("sup"):
                    r = t.text_frame.paragraphs[0].add_run()
                    r.text = box["sup"]
                    r.font._rPr.set("baseline", "30000")
            if spec.get("chart"):
                ch = spec["chart"]
                data = new_chart_data()
                data.categories = ch["categories"]
                for sname, vals in ch["series"]:
                    data.add_series(sname, vals)
                kind = {
                    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
                    "pie": XL_CHART_TYPE.PIE,
                    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                }[ch["type"]]
                gf = s.shapes.add_chart(kind, Inches(7), Inches(2), Inches(5.5), Inches(4), data)
                if ch.get("invisible_first"):
                    gf.chart.plots[0].series[0].format.fill.background()
                if ch.get("white_first"):  # the deck-kit hides a base series by filling it white
                    from pptx.dml.color import RGBColor

                    make_rgb: Callable[[int, int, int], RGBColor] = RGBColor  # unannotated in python-pptx
                    fill = gf.chart.plots[0].series[0].format.fill
                    fill.solid()
                    fill.fore_color.rgb = make_rgb(0xFF, 0xFF, 0xFF)
            if spec.get("table"):
                rows = spec["table"]
                gt = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(2), Inches(6), Inches(3))
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        gt.table.cell(r, c).text = val
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = spec["notes"]
        out = tmp_path / name
        prs.save(str(out))
        return out

    return build
