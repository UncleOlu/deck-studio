import json
from pathlib import Path as _P2

from conftest import load_script

storyline = load_script("storyline-lint")
trace = load_script("trace-check")
integrity = load_script("integrity-check")
layout = load_script("layout-check")
import deck_model  # noqa: E402

GOOD = "Labour hours do not flex with traffic, so 11% of store hours are idle"


def test_storyline_flags_topic_label_question_and_missing_source(make_pptx):
    deck = make_pptx([
        {"title": "Harvest Lane Markets board review"},
        {"title": "Executive summary", "body": "Three findings follow."},
        {"title": "Labour overview", "body": "Stores schedule labour by hand in spreadsheets today."},
        {"title": "Why is shrink so high?"},                             # question
        {"title": GOOD, "chart": {"type": "bar", "categories": ["A", "B"], "series": [("s", (1, 2))]}},
    ])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    text = "\n".join(errors)
    assert "slide 3: topic label" in text
    assert "slide 4: title is a question" in text
    assert "slide 5: chart/table slide has no 'Source:'" in text


def test_storyline_passes_a_clean_consulting_deck(make_pptx):
    deck = make_pptx([
        {"title": "Harvest Lane Markets board review"},
        {"title": "Executive summary: labour and shrink can add 150 bps of EBIT margin", "body": "a"},
        {"title": GOOD, "body": "Source: store P&L FY25",
         "chart": {"type": "bar", "categories": ["A", "B"], "series": [("s", (1, 2))]}},
    ])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    assert errors == []


def test_storyline_requires_early_exec_summary(make_pptx):
    deck = make_pptx([{"title": "Cover page"}] + [{"title": GOOD}] * 3 +
                     [{"title": "Executive summary: we recommend a two-wave labour programme"}])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    assert any("executive summary must sit at or before slide 3" in e for e in errors)


def write_facts(tmp_path, rows):
    p = tmp_path / "facts.jsonl"
    p.write_text("\n".join(json.dumps(r) for r in rows) + "\n")
    return p


FACT = {"id": "F0027", "value": 4820.0, "unit": "USD", "scale": 1e6, "period": "FY25",
        "text": "Revenue", "file": "store_pnl.xlsx", "locator": "Summary!B3"}


def test_trace_passes_cited_rescaled_and_derived_numbers(make_pptx, tmp_path):
    facts = write_facts(tmp_path, [FACT])
    deck = make_pptx([{"title": "Revenue of $4.8B funds the programme", "body": "Each 10 bps is worth $4.8M",
                       "notes": "Revenue [F0027]. calc: $4.8M = 10 bps x $4.82B [F0027]"}])
    problems, total = trace.check(str(deck), str(facts), None)
    assert total == 3 and problems == []  # $4.8B, 10 bps, $4.8M


def test_trace_flags_untraced_and_unknown_ids(make_pptx, tmp_path):
    facts = write_facts(tmp_path, [FACT])
    deck = make_pptx([{"title": "Revenue of $4.8B funds a $72M programme", "notes": "[F0027] [F9999]"}])
    problems, _ = trace.check(str(deck), str(facts), None)
    assert any("'$72M'" in p for p in problems)
    assert any("F9999" in p for p in problems)


def test_trace_accepts_illustrative_data_contract(make_pptx, tmp_path):
    facts = write_facts(tmp_path, [FACT])
    data = tmp_path / "deck-data.json"
    data.write_text(json.dumps({"illustrative": True, "programme_cost_m": 72}))
    deck = make_pptx([{"title": "A $72M programme pays back in two years", "notes": "illustrative"}])
    problems, _ = trace.check(str(deck), str(facts), str(data))
    assert problems == []


def test_integrity_catches_broken_waterfall_and_passes_good_one(make_pptx):
    good = {"type": "stacked", "categories": ["FY24", "Labour", "Shrink", "FY26"], "invisible_first": True,
            "series": [("base", (0, 4.1, 4.9, 0)), ("value", (4.1, 0.8, 0.7, 5.6))]}
    bad = {"type": "stacked", "categories": ["FY24", "Labour", "Shrink", "FY26"], "invisible_first": True,
           "series": [("base", (0, 4.1, 5.2, 0)), ("value", (4.1, 0.8, 0.7, 5.6))]}
    ok_deck = deck_model.load(make_pptx([{"title": GOOD, "chart": good}], "ok.pptx"))
    bad_deck = deck_model.load(make_pptx([{"title": GOOD, "chart": bad}], "bad.pptx"))
    assert integrity.check_waterfall(1, ok_deck.slides[0].charts[0]) == []
    errs = integrity.check_waterfall(1, bad_deck.slides[0].charts[0])
    assert errs and "Shrink" in errs[0]


def test_integrity_pie_table_consistency_footnotes(make_pptx):
    deck = deck_model.load(make_pptx([
        {"title": "Fresh drives 62% of shrink across the estate",
         "chart": {"type": "pie", "categories": ["Fresh", "Other"], "series": [("share", (62, 36))]},
         "table": [["Format", "Stores"], ["Superstore", "40"], ["Express", "32"], ["Total", "142"]],
         "boxes": [{"x": 0.6, "y": 6.6, "w": 6, "h": 0.4, "text": "Shrink rate", "sup": "2", "pt": 10}]},
        {"title": "Fresh drives 58% of shrink in Express stores"},
    ]))
    errs = []
    for s in deck.slides:
        for ch in s.charts:
            errs += integrity.check_shares(s.index, ch)
        for t in s.tables:
            errs += integrity.check_table(s.index, t.rows)
        errs += integrity.check_footnotes(s)
    text = "\n".join(errs)
    assert "sum to 98%" in text
    assert "parts sum to 72" in text
    assert "footnote markers appear as 2" in text


def test_integrity_one_value_per_metric():
    deck = deck_model.Deck(path=None, kind="pptx", slides=[
        deck_model.Slide(1, texts=["t", "Adjusted EBITDA of $91M in FY25"]),
        deck_model.Slide(2, texts=["t", "Adjusted EBITDA of $86M in FY25"]),
        deck_model.Slide(3, texts=["t", "94% net revenue retention vs. 87% benchmark"]),
    ])
    errs = integrity.check_consistency(deck)
    assert len(errs) == 1 and "adjusted, ebitda" in errs[0]


def test_layout_flags_overflow_drift_and_near_miss(make_pptx):
    long = "word " * 120
    deck = deck_model.load(make_pptx([
        {"title": "Cover"},
        {"title": GOOD, "boxes": [{"x": 0.6, "y": 2, "w": 3, "h": 0.6, "text": long, "name": "Crowded"}]},
        {"title": GOOD, "title_y": 0.5},
        {"title": GOOD, "title_y": 0.9, "boxes": [{"x": 0.66, "y": 3, "w": 3, "h": 1, "text": "near miss"}]},
        {"title": GOOD},
        {"title": "Close"},
    ]))
    errs, warns = layout.geometry_checks(deck)
    text = "\n".join(errs + warns)
    assert "'Crowded' text needs" in text
    assert "slide 4: title y is 0.90" in text
    assert "nearly align" in text


validator = load_script("validate_pptx")


def _rewrite(src, dst, edit):
    import zipfile
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = edit(item.filename, zin.read(item.filename))
            if data is not None:
                zout.writestr(item, data)


def test_validate_pptx_passes_clean_and_catches_planted_defects(make_pptx, tmp_path):
    good = make_pptx([{"title": GOOD, "body": "x"}, {"title": GOOD}])
    assert validator.check(good) == ([], [])

    def hash_colour(name, data):
        if name == "ppt/slides/slide1.xml":
            return data.replace(b"<a:t>", b'<a:solidFill><a:srgbClr val="#FF0000"/></a:solidFill><a:t>', 0)
        return data

    def drop_slide(name, data):
        return None if name == "ppt/slides/slide2.xml" else data

    bad_xml = tmp_path / "badxml.pptx"
    _rewrite(good, bad_xml, lambda n, d: d.replace(b"</p:sld>", b"<oops></p:sld>") if n.endswith("slide1.xml") else d)
    missing = tmp_path / "missing.pptx"
    _rewrite(good, missing, drop_slide)
    colour = tmp_path / "colour.pptx"
    _rewrite(good, colour, lambda n, d: d.replace(b'<a:srgbClr val="', b'<a:srgbClr val="#', 1)
             if n == "ppt/theme/theme1.xml" else d)
    assert any("not well-formed" in e for e in validator.check(bad_xml)[0])
    assert any("missing part" in e or "no slide part" in e for e in validator.check(missing)[0])
    assert any("6-digit hex" in e for e in validator.check(colour)[0])


def test_consistency_ignores_ranges_modifiers_and_pages():
    deck = deck_model.Deck(path=None, kind="pptx", slides=[
        deck_model.Slide(1, texts=["t", "Peers trade at 9.0x–11.0x LTM EBITDA, median 10.7x", "Confidential draft 2"]),
        deck_model.Slide(2, texts=["t", "Strategic buyers paid 11.8x–12.5x LTM EBITDA", "Confidential draft 3"]),
        deck_model.Slide(3, texts=["t", "premium to the 52-week average of $31.37"]),
        deck_model.Slide(4, texts=["t", "26-week average $32.89"]),
    ])
    assert integrity.check_consistency(deck) == []


def test_banking_labels_pass_but_long_banking_claims_use_consulting_limit(make_pptx):
    deck = make_pptx([
        {"title": "Project Ash discussion materials"},
        {"title": "Summary of analyses", "body": "($ in millions)\nSource: FactSet"},
        {"title": "Selected precedent transactions", "body": "Source: company filings",
         "table": [["Date", "Target", "EV/EBITDA"], ["2024", "A", "10.2x"], ["2025", "B", "11.8x"]]},
        {"title": "Strategic buyers can pay more than sponsors because the cost synergies they gain from the "
                  "combination are larger than any sponsor's operating plan"},
    ])
    errors, warns, _ = storyline.lint(str(deck), "banking")
    assert not any("slide 3" in e for e in errors)                    # neutral label, sourced
    assert any("slide 4" in e and "limit is 20" in e for e in errors)  # a claim: held to the consulting limit
    assert any("slide 3" in w and "unit line" in w for w in warns)


def test_answer_title_with_three_claims_counts_as_exec_summary(make_pptx):
    body = ("Labour scheduling can release 90 bps within 12 months\n"
            "Fresh markdown timing can release 60 bps within 9 months\n"
            "Self-checkout can release 20 bps once staff turnover falls below 40%")
    deck = make_pptx([{"title": "Harvest Lane board review"},
                      {"title": "Labour and shrink can add 170 bps of EBIT margin by FY27", "body": body},
                      {"title": GOOD}])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    assert not any("executive summary" in e for e in errors)


def test_sample_deck_json_match_their_schemas():
    import json as _json
    from pathlib import Path as _P

    from jsonschema import Draft202012Validator
    from referencing import Registry, Resource

    tdir = _P(__file__).resolve().parents[1] / "skills" / "premium-decks" / "templates"
    kit = _json.loads((tdir / "deck-kit.schema.json").read_text())
    registry = Registry().with_resource("deck-kit.schema.json", Resource.from_contents(kit))
    for schema_name, sample in (("consulting-deck.schema.json", "consulting-deck.sample.json"),
                                ("pitchbook.schema.json", "pitchbook.sample.json"),
                                ("pitchbook.schema.json", "pitchbook-pitch.sample.json")):
        schema = _json.loads((tdir / schema_name).read_text())
        errors = list(Draft202012Validator(schema, registry=registry).iter_errors(
            _json.loads((tdir / sample).read_text())))
        assert not errors, f"{sample}: {errors[0].message} at {list(errors[0].absolute_path)}"


def test_integrity_flags_truncated_bar_axis():
    series = [{"name": "s", "values": [12.4, 13.9], "invisible": False}]
    ch = deck_model.Chart("barChart", "clustered", ["A", "B"], series, axis_min=11.5)
    assert integrity.check_bar_baseline(1, ch)
    ch.axis_min = 0
    assert integrity.check_bar_baseline(1, ch) == []
    ch.axis_min = None  # automatic axis with close values: PowerPoint would truncate
    assert integrity.check_bar_baseline(1, ch)


def test_integrity_checks_kit_style_white_base_waterfall(make_pptx):
    bad = {"type": "stacked", "categories": ["FY24", "Labour", "Shrink", "FY26"], "white_first": True,
           "series": [("base", (0, 4.1, 5.2, 0)), ("value", (4.1, 0.8, 0.7, 5.6))]}
    deck = deck_model.load(make_pptx([{"title": GOOD, "chart": bad}], "white.pptx"))
    ch = deck.slides[0].charts[0]
    assert ch.series[0]["invisible"] and ch.bar_dir == "col"
    assert integrity.check_waterfall(1, ch)


def test_validate_pptx_flags_duplicate_shape_ids(make_pptx, tmp_path):
    good = make_pptx([{"title": GOOD, "body": "x"}])
    dup = tmp_path / "dup.pptx"
    _rewrite(good, dup,
             lambda n, d: d.replace(b'<p:cNvPr id="3"', b'<p:cNvPr id="2"') if n.endswith("slide1.xml") else d)
    assert any("duplicate shape ids" in e for e in validator.check(dup)[0])


def test_storyline_flags_gap_pages_and_input_talk_but_allows_them_in_appendix(make_pptx):
    deck = make_pptx([
        {"title": "Project Ash discussion materials"},
        {"title": "Executive summary: the offer sits above four of five valuation ranges", "body": "a"},
        {"title": "Analyses pending further information", "body": "LBO analysis: not performed."},
        {"title": GOOD, "body": "Note: the file stores the formula without a value.\nSource: store P&L FY25"},
        {"title": GOOD, "body": "Note: EBIT recomputed from the Summary sheet.\nSource: store P&L FY25"},
        {"title": "Appendix"},
        {"title": "Data used and open items", "body": "Premiums data was not provided."},
    ])
    errors, warns, _ = storyline.lint(str(deck), "consulting")
    text = "\n".join(errors)
    assert "slide 3: data-gap page" in text
    assert "slide 4: slide text talks about the inputs" in text
    assert "slide 5" not in text                        # a business-source note passes
    assert not any("slide 7" in e for e in errors)      # the appendix may list gaps
    assert any("slide 7" in w for w in warns)           # but still names the source, not "the data"
    kerrors, _, _ = storyline.lint(str(deck), "keynote")
    assert not any("inputs" in e or "data-gap" in e for e in kerrors)


def test_pptx2pdf_passes_paths_as_argv_not_script_text(tmp_path, monkeypatch):
    import subprocess
    import sys
    p2p = load_script("pptx2pdf")
    seen = {}

    def fake_run(argv, **kw):
        seen["argv"] = argv
        return subprocess.CompletedProcess(argv, 0, b"deck", b"")

    monkeypatch.setattr(p2p.subprocess, "run", fake_run)
    monkeypatch.setattr(p2p.os.path, "exists", lambda _p: True)
    monkeypatch.setattr(sys, "platform", "darwin")
    evil = str(tmp_path / 'x" & do shell script "touch /tmp/pwned" & ".pptx')
    p2p.applescript_convert("Microsoft PowerPoint", evil, str(tmp_path / "out.pdf"))
    script = seen["argv"][2]
    assert "do shell script" not in script and evil in seen["argv"][3:]


def test_storyline_no_false_positives_on_steerco_pending_data_room_and_marked_unit_line(make_pptx):
    deck = make_pptx([
        {"title": "Project Cobalt discussion materials"},
        {"title": "Executive summary: the offer sits above four of five valuation ranges", "body": "a"},
        {"title": "Two decisions pending: approve the finance go-live date and release the budget"},
        {"title": "Process: 14 buyers enter the data room in week 3 and bid by week 8"},
        {"title": "Trading comps: the offer is 12.5x LTM EBITDA against a 10.0x median",
         "body": "(Implied value per share, $) (1)\nSource: FactSet as of 15 May 2026",
         "table": [["Peer", "EV/EBITDA"], ["A", "10.0x"], ["B", "9.2x"]]},
        {"title": "Precedents: the offer is 12.5x against a 10.9x transaction median",
         "body": "Source: precedents file", "table": [["Deal", "x"], ["A", "9.8x"], ["B", "11.9x"]]},
    ])
    errors, warns, _ = storyline.lint(str(deck), "banking")
    text = "\n".join(errors + warns)
    assert "slide 3" not in text and "slide 4" not in text
    assert not any("slide 5" in w and "unit line" in w for w in warns)
    assert any("slide 6" in e and "inputs" in e for e in errors)    # "precedents file" names the build, not a source


def test_trace_section_strip_skip_does_not_hide_data_lines(make_pptx, tmp_path):
    facts = tmp_path / "facts.jsonl"
    facts.write_text("")
    deck = make_pptx([
        {"title": "Cover"},
        {"title": GOOD, "body": "1 Baseline · 2 Labour · 3 Shrink\nPilot Stores 240 · Control Stores 238"},
    ])
    code, out = subprocess_run_trace(deck, facts)
    assert "'240'" in out and "'238'" in out          # data lines are still checked
    assert "'2'" not in out and "'3'" not in out      # the section strip is not


def subprocess_run_trace(deck, facts):
    import subprocess
    import sys
    from pathlib import Path
    script = Path(__file__).resolve().parents[1] / "skills" / "premium-decks" / "scripts" / "trace-check.py"
    r = subprocess.run([sys.executable, str(script), str(deck), str(facts)], capture_output=True, text=True)
    return r.returncode, r.stdout


def test_storyline_soft_gap_titles_need_missing_data_evidence(make_pptx):
    deck = make_pptx([
        {"title": "Project Birch discussion materials"},
        {"title": "Executive summary: the offer sits above four of five valuation ranges", "body": "a"},
        {"title": "Open items", "body": "Store-level labour data not yet received; awaiting Q3 figures."},
        {"title": "Synergy case to be confirmed", "body": "Procurement savings: TBC pending data."},
        {"title": "Open items: two decisions for the SteerCo by 30 October",
         "body": "Approve the go-live date\nRelease the phase 2 budget"},
    ])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    text = "\n".join(errors)
    assert "slide 3: data-gap page" in text
    assert "slide 4: data-gap page" in text
    assert "slide 5" not in text                        # decision lists stay legitimate


def test_storyline_flags_any_named_input_file_as_a_source(make_pptx):
    table = [["Line", "FY25"], ["Opex", "12.0"], ["Capex", "3.1"]]
    deck = make_pptx([
        {"title": "Project Birch discussion materials"},
        {"title": "Executive summary: the offer sits above four of five valuation ranges", "body": "a"},
        {"title": "Opex fell 8% to $12.0m while capex held at $3.1m",
         "body": "Source: budget tracker file", "table": table},
        {"title": "Opex fell 8% to $12.0m while capex held at $3.1m",
         "body": "Source: headcount spreadsheet", "table": table},
        {"title": "Opex fell 8% to $12.0m while capex held at $3.1m",
         "body": "Source: company 10-K filing; HR master file", "table": table},
    ])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    assert any("slide 3" in e and "inputs" in e for e in errors)
    assert any("slide 4" in e and "inputs" in e for e in errors)
    assert not any("slide 5" in e and "inputs" in e for e in errors)


def test_integrity_checks_inline_parenthesised_footnote_markers(make_pptx):
    bad = make_pptx([{"title": "Comps: the offer is 12.5x against a 10.0x median",
                      "body": "Revenue (2) grew 4%\nEBITDA (1) margin 18%\n(1) Adjusted\n(2) Reported"}],
                    name="bad.pptx")
    ok = make_pptx([{"title": "Comps: the offer is 12.5x against a 10.0x median",
                     "body": "Revenue (1) grew 4%\nEBITDA (2) fell $(3)m\n(1) Adjusted\n(2) Reported"}],
                   name="ok.pptx")
    missing = make_pptx([{"title": "Comps: the offer is 12.5x against a 10.0x median",
                          "body": "Revenue (1) grew 4%\nEBITDA (2) fell\n(1) Adjusted\nSource: FactSet"}],
                        name="missing.pptx")
    def run(path):
        return "\n".join(integrity.check_footnotes(deck_model.load(str(path)).slides[0]))
    assert "footnote markers appear as 2, 1" in run(bad)
    assert run(ok) == ""
    assert "marker 2 has no note line" in run(missing)


def test_deck_edit_dup_delete_move_keep_the_package_valid(make_pptx, tmp_path):
    from pptx import Presentation
    edit = load_script("deck_edit")
    validator = load_script("validate_pptx")
    chart = {"type": "bar", "categories": ["A", "B"], "series": [("S", (1, 2))]}
    deck = make_pptx([{"title": "One", "notes": "speaker one"}, {"title": "Two", "chart": chart},
                      {"title": "Three"}])
    out = tmp_path / "out.pptx"
    assert edit.main([str(deck), "-o", str(out), "dup", "2", "--after", "3"]) == 0
    prs = Presentation(str(out))
    assert [line.split("]  ")[1] for line in edit.listing(out)] == ["One", "Two", "Three", "Two"]
    charts = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 2 and charts[0].part is not charts[1].part     # the copy edits independently
    assert tuple(charts[1].plots[0].series[0].values) == (1, 2)
    assert charts[1].part.chart_workbook.xlsx_part is not charts[0].part.chart_workbook.xlsx_part
    assert validator.check(str(out))[0] == []
    assert edit.main([str(out), "move", "4", "--to", "1"]) == 0
    assert edit.main([str(out), "delete", "3"]) == 0
    prs = Presentation(str(out))
    assert [line.split("]  ")[1] for line in edit.listing(out)] == ["Two", "One", "Three"]
    assert validator.check(str(out))[0] == []
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart) == 1
    import zipfile
    names = zipfile.ZipFile(out).namelist()
    assert sum(n.startswith("ppt/charts/chart") for n in names) == 1       # no orphaned chart part


def test_doctor_reports_required_items_and_finds_soffice_off_path(monkeypatch):
    doctor = load_script("doctor")
    rows = {item: (ok, required) for item, ok, required, _ in doctor.check()}
    assert rows["python-pptx"] == (True, True) and rows["pptxgenjs (vendored)"] == (True, True)
    p2p = load_script("pptx2pdf")
    monkeypatch.setattr(p2p.shutil, "which", lambda _n: None)
    monkeypatch.setattr(p2p.os.path, "isfile", lambda c: c == p2p.SOFFICE_CANDIDATES[1])
    assert p2p.find_soffice() == p2p.SOFFICE_CANDIDATES[1]


def test_storyline_accepts_key_considerations_as_the_pitch_summary_page(make_pptx):
    def errors_for(summary_title):
        deck = make_pptx([
            {"title": "Project Wren discussion materials"},
            {"title": "Situation overview"},
            {"title": summary_title, "body": "a"},
            {"title": "Share price performance", "body": "Source: FactSet as of 1 Sep 2026",
             "table": [["Date", "Price"], ["Jan", "10.0"], ["Jun", "12.0"]]},
        ], name=f"{summary_title}.pptx")
        return "\n".join(storyline.lint(str(deck), "banking")[0])
    assert "executive summary must sit" not in errors_for("Key considerations")
    assert "executive summary must sit" in errors_for("Company background")


def test_integrity_reads_html_waterfall_custom_properties(tmp_path):
    def deck(mid_from, mid_to):
        html = ('<section class="slide"><h2>Bridge</h2><div class="wf" style="--min:0;--max:45">'
                '<div class="wf-bar" data-label="FY25" style="--from:0;--to:36.8"><i></i></div>'
                f'<div class="wf-bar" data-label="Step" style="--from:{mid_from};--to:{mid_to}"><i></i></div>'
                '<div class="wf-bar" data-label="End" style="--from:0;--to:39.2"><i></i></div></div></section>')
        p = tmp_path / f"wf-{mid_from}.html"
        p.write_text(html)
        chart = deck_model.load(str(p)).slides[0].charts[0]
        return integrity.check_waterfall(1, chart)
    assert deck(36.8, 39.2) == []
    assert any("floats at 30" in e for e in deck(30, 32.4))


def test_kit_refuses_one_callout_device_on_most_slides(tmp_path):
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        import pytest
        pytest.skip("node not installed")
    from pathlib import Path as _P
    builder = _P(__file__).resolve().parents[1] / "skills" / "premium-decks" / "templates" / "build-consulting.js"
    side = {"points": [{"lead": "Gap.", "text": "Two points above peers."}]}

    def spec(n_side):
        slides = [{"type": "cover", "title": "Project Birch"},
                  {"type": "exec-summary", "title": "Opex can fall 8% by FY27",
                   "points": [{"lead": "a"}, {"lead": "b"}], "notes": "n"}]
        for i in range(6):
            sl = {"type": "chart", "title": f"Opex line {i} fell 8% in FY25", "labels": ["FY24", "FY25"],
                  "series": [{"name": "Opex", "values": [10, 9.2]}], "source": "Company P&L FY25", "notes": "n"}
            if i < n_side:
                sl["side"] = side
            slides.append(sl)
        path = tmp_path / f"deck-{n_side}.json"
        path.write_text(json.dumps({"output": f"out-{n_side}.pptx", "slides": slides}))
        return subprocess.run([node, str(builder), str(path)], capture_output=True, text=True, cwd=tmp_path,
                              timeout=120)
    assert spec(2).returncode == 0
    bad = spec(5)
    assert bad.returncode != 0 and "device mix" in bad.stderr


def test_lint_rules_do_not_fire_on_ordinary_business_text(make_pptx):
    table = [["Line", "FY25"], ["Opex", "12.0"], ["Capex", "3.1"]]
    deck = make_pptx([
        {"title": "Project Birch discussion materials"},
        {"title": "Executive summary: the offer sits above four of five valuation ranges", "body": "a"},
        {"title": "Open items", "body": "Awaiting CFO sign-off on the TSA\nTSA exit dates TBC"},
        {"title": "Opex fell 8% to $12.0m after Halden files for Chapter 11 protection",
         "body": "Source: court filings; the company will file its 10-K in March", "table": table},
    ])
    errors, _, _ = storyline.lint(str(deck), "consulting")
    assert not any("slide 3" in e and "data-gap" in e for e in errors)
    assert not any("slide 4" in e and "inputs" in e for e in errors)


def test_inline_markers_ignore_numbered_labels_and_bracketed_negatives(make_pptx):
    plain = make_pptx([{"title": "Comps: the offer is 12.5x against a 10.0x median",
                        "body": "Phase (2) starts in May\nNet debt (12) at close"}], name="plain.pptx")
    assert integrity.check_footnotes(deck_model.load(str(plain)).slides[0]) == []


def test_integrity_main_checks_html_waterfalls_and_flags_unreadable_bars(tmp_path):
    import subprocess
    import sys
    script = _P2(__file__).resolve().parents[1] / "skills" / "premium-decks" / "scripts" / "integrity-check.py"
    good = ('<section class="slide"><h2>Bridge</h2><div class="wf" style="--min:0;--max:45">'
            '<div class="wf-bar" data-label="A" style="--from:0;--to:36.8"></div>'
            '<div class="wf-bar" data-label="B" style="--from:36.8;--to:39.2"></div>'
            '<div class="wf-bar" data-label="C" style="--from:0;--to:39.2"></div></div></section>')
    bad = good.replace("--from:36.8;--to:39.2", "--from:calc(1 + 2);--to:.")
    for name, html, code in (("good", good, 0), ("bad", bad, 1)):
        p = tmp_path / f"{name}.html"
        p.write_text(html)
        r = subprocess.run([sys.executable, str(script), str(p)], capture_output=True, text=True)
        assert r.returncode == code, r.stdout + r.stderr
        if name == "bad":
            assert "cannot read" in r.stdout

