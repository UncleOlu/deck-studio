#!/usr/bin/env python3
"""Validate a .pptx package: structure, relationships, XML, and known corrupters.

Checks (each failure names its fix):
  1. The file is a readable zip with the three required parts.
  2. Every part has a content type ([Content_Types].xml Default or Override).
  3. Every internal relationship target exists; relationship ids are unique.
  4. Every XML part is well-formed (parsed with defusedxml).
  5. Slide ids in presentation.xml are unique and >= 256; every listed slide exists.
  6. Colour values are 6-digit hex (no '#', no alpha digits).
  6b. Shape ids are unique within each slide (duplicates make PowerPoint offer a repair).
  7. Stacked bar/column charts do not use data labels at 'outEnd' (PowerPoint
     rejects the file).
  8. python-pptx can open the package and save it again.
  9. Optional --render: the deck converts to PDF (LibreOffice, PowerPoint, or Keynote).
 10. Optional --native-charts: no picture is used where a chart belongs (warns
     when a content slide has pictures but no chart or table).

Exit 0 when valid, 1 otherwise. Usage: python3 validate_pptx.py deck.pptx [--render] [--native-charts]
"""

from __future__ import annotations

import argparse
import contextlib
import io
import posixpath
import re
import sys
import zipfile
from pathlib import Path

from defusedxml import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pptx2pdf  # noqa: E402

CT_NS = "{http://schemas.openxmlformats.org/package/2006/content-types}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
REQUIRED = ("[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml")
HEX6 = re.compile(r"^[0-9A-Fa-f]{6}$")


def rels_base(rels_part: str) -> str:
    # ppt/slides/_rels/slide1.xml.rels -> ppt/slides
    folder = posixpath.dirname(posixpath.dirname(rels_part))
    return folder


def check(path: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        return [f"{path.name} is not a zip archive — rebuild it; do not rename other formats to .pptx"], []
    with z:
        bad = z.testzip()
        if bad:
            errors.append(f"corrupt zip member {bad} — rebuild the file")
        names = set(z.namelist())
        for req in REQUIRED:
            if req not in names:
                errors.append(f"missing required part {req} — the generator did not write a complete package")
        if errors:
            return errors, warnings

        ct = ET.fromstring(z.read("[Content_Types].xml"))
        defaults = {d.get("Extension", "").lower() for d in ct.findall(f"{CT_NS}Default")}
        overrides = {o.get("PartName", "").lstrip("/") for o in ct.findall(f"{CT_NS}Override")}
        for n in names:
            if n.endswith("/") or n == "[Content_Types].xml":
                continue
            ext = n.rsplit(".", 1)[-1].lower() if "." in n else ""
            if n not in overrides and ext not in defaults:
                errors.append(f"{n} has no content type — add an Override or a Default for .{ext}")

        for n in sorted(names):
            if n.endswith(".xml") or n.endswith(".rels"):
                try:
                    root = ET.fromstring(z.read(n))
                except ET.ParseError as err:
                    errors.append(f"{n} is not well-formed XML ({err}) — escape &, <, > in text")
                    continue
                if n.endswith(".rels"):
                    ids = [r.get("Id") for r in root.findall(f"{REL_NS}Relationship")]
                    dup = {i for i in ids if ids.count(i) > 1}
                    if dup:
                        errors.append(f"{n}: duplicate relationship ids {sorted(dup)}")
                    base = rels_base(n)
                    for r in root.findall(f"{REL_NS}Relationship"):
                        if r.get("TargetMode") == "External":
                            continue
                        target = r.get("Target", "")
                        full = (
                            target.lstrip("/")
                            if target.startswith("/")
                            else posixpath.normpath(posixpath.join(base, target))
                        )
                        if full not in names:
                            errors.append(f"{n}: relationship {r.get('Id')} points to missing part {full}")
                if re.match(r"ppt/slides/slide\d+\.xml$", n):
                    ids = [
                        el.get("id")
                        for el in root.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
                    ]
                    dups = sorted({i for i in ids if ids.count(i) > 1})
                    if dups:
                        errors.append(
                            f"{n}: duplicate shape ids {dups} — PowerPoint may ask to repair the file; "
                            "give every shape on a slide a unique id"
                        )
                for el in root.iter():
                    tag = el.tag.rsplit("}", 1)[-1]
                    if tag == "srgbClr" and not HEX6.match(el.get("val", "")):
                        errors.append(f"{n}: colour '{el.get('val')}' is not 6-digit hex — drop '#' and alpha digits")
                if n.startswith("ppt/charts/chart") and n.endswith(".xml"):
                    for bar in root.iter(f"{C_NS}barChart"):
                        grouping = bar.find(f"{C_NS}grouping")
                        if grouping is not None and grouping.get("val") in ("stacked", "percentStacked"):
                            for pos in bar.iter(f"{C_NS}dLblPos"):
                                if pos.get("val") == "outEnd":
                                    errors.append(
                                        f"{n}: stacked bar uses data labels at outEnd — use ctr, inEnd, or inBase"
                                    )

        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        prels = (
            ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
            if "ppt/_rels/presentation.xml.rels" in names
            else None
        )
        targets = (
            {r.get("Id"): r.get("Target") for r in prels.findall(f"{REL_NS}Relationship")} if prels is not None else {}
        )
        lst = pres.find(f"{P_NS}sldIdLst")
        seen = set()
        for sid in lst.findall(f"{P_NS}sldId") if lst is not None else []:
            num = int(sid.get("id", "0"))
            if num < 256 or num in seen:
                errors.append(f"presentation.xml: slide id {num} is duplicate or below 256")
            seen.add(num)
            tgt = targets.get(sid.get(R_ID))
            if not tgt or posixpath.normpath(posixpath.join("ppt", tgt)) not in names:
                errors.append(f"presentation.xml: listed slide {sid.get(R_ID)} has no slide part")

    try:
        from pptx import Presentation

        prs = Presentation(str(path))
        prs.save(io.BytesIO())
    except Exception as err:  # noqa: BLE001 — any failure here means PowerPoint may reject the file too
        errors.append(f"python-pptx cannot round-trip the file ({type(err).__name__}: {err})")
    return errors, warnings


def native_chart_warnings(path: Path) -> list[str]:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import deck_model

    out = []
    for s in deck_model.load(path).slides[1:]:
        pics = [x for x in s.shapes if x.kind == "picture"]
        if pics and not s.charts and not s.tables and any(x.w * x.h > 6 for x in pics):
            out.append(
                f"slide {s.index}: a large picture and no native chart or table — "
                "if it shows data, rebuild it as an editable chart"
            )
        for x in pics:
            if not x.alt.strip():
                out.append(f"slide {s.index}: picture '{x.name}' has no alt text")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="Validate a .pptx package.")
    ap.add_argument("deck")
    ap.add_argument("--render", action="store_true", help="also convert to PDF with scripts/pptx2pdf.py")
    ap.add_argument("--native-charts", action="store_true", help="warn on pictures used where charts belong")
    args = ap.parse_args()
    path = Path(args.deck)
    errors, warnings = check(path)
    if args.native_charts and not errors:
        warnings += native_chart_warnings(path)
    if args.render and not errors:
        log = io.StringIO()
        with contextlib.redirect_stdout(log):
            rc = pptx2pdf.convert(str(path))
        if rc != 0:
            lines = log.getvalue().strip().splitlines()
            errors.append("render failed: " + (lines[-1] if lines else f"exit {rc}"))
    for e in errors:
        print(f"ERROR {e}")
    for w in warnings:
        print(f"WARN {w}")
    print(f"validate_pptx: {'PASSED' if not errors else 'FAILED'} — {len(errors)} error(s), {len(warnings)} warning(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
