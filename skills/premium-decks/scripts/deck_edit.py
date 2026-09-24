#!/usr/bin/env python3
"""Structural edits to an existing .pptx: list, duplicate, delete, and move slides.

Use it for template work before editing slide content. Each command writes a
new file (-o) or rewrites the input in place. python-pptx saves only the
parts that are still referenced, so a deleted slide's media, charts, and notes
go with it (unless another slide still links to it). Transitions, animations,
and rich-text notes are not copied by `dup`; notes are copied as plain text.

  list                        print slide number, layout, and title
  dup N [--after M]           copy slide N (charts and their workbooks are
                              copied too, so the copy edits independently);
                              placed after M, default after N
  delete N [N ...]            remove slides
  move N --to M               move slide N to position M

Slide numbers are 1-based. Usage:
  python3 deck_edit.py deck.pptx list
  python3 deck_edit.py deck.pptx dup 3 --after 5 -o out.pptx
  python3 deck_edit.py deck.pptx delete 2 7
  python3 deck_edit.py deck.pptx move 4 --to 2
"""

from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart

SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}


def _slide_ids(prs: Any) -> Any:
    return prs.slides._sldIdLst


def _check(prs: Any, *numbers: int) -> None:
    n = len(prs.slides)
    for k in numbers:
        if not 1 <= k <= n:
            sys.exit(f"deck_edit: slide {k} is out of range (deck has {n} slides)")


R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _remap_rids(root: Any, rid_map: dict[str, str]) -> None:
    for el in root.iter():
        for attr, val in list(el.attrib.items()):
            if attr.startswith(f"{{{R_NS}}}") and val in rid_map:
                el.set(attr, rid_map[val])


def _copy_chart(src_part: ChartPart, package: Any) -> ChartPart:
    """A new chart part with its own embedded workbook; other chart relationships
    (user-shape drawings, links) point at the same targets as the source."""
    partname = package.next_partname(ChartPart.partname_template)
    new: ChartPart = ChartPart.load(partname, src_part.content_type, package, src_part.blob)
    rid_map: dict[str, str] = {}
    for rid, rel in src_part.rels.items():
        if rel.is_external:
            rid_map[rid] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.PACKAGE:
            rid_map[rid] = new.relate_to(EmbeddedXlsxPart.new(rel.target_part.blob, package), rel.reltype)
        else:
            rid_map[rid] = new.relate_to(rel.target_part, rel.reltype)
    _remap_rids(new._element, rid_map)
    return new


def duplicate(prs: Any, n: int, after: int | None) -> int:
    _check(prs, n)
    src = prs.slides[n - 1]
    dst = prs.slides.add_slide(src.slide_layout)
    for shape in list(dst.shapes):  # drop the layout's placeholders; the copy brings its own
        shape._element.getparent().remove(shape._element)
    rid_map: dict[str, str] = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in SKIP_RELS:
            continue
        if rel.is_external:
            rid_map[rid] = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif isinstance(rel.target_part, ChartPart):
            rid_map[rid] = dst.part.relate_to(_copy_chart(rel.target_part, prs.part.package), rel.reltype)
        else:  # images and media are shared read-only between slides
            rid_map[rid] = dst.part.relate_to(rel.target_part, rel.reltype)
    tree = dst.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    _remap_rids(tree, rid_map)
    src_bg = src._element.cSld.bg
    if src_bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(src_bg))
    if src.has_notes_slide and src.notes_slide.notes_text_frame is not None:
        dst.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    target = (after if after is not None else n) + 1
    move(prs, len(prs.slides), target)
    return target


def delete(prs: Any, numbers: list[int]) -> None:
    _check(prs, *numbers)
    ids = _slide_ids(prs)
    entries = list(ids)
    for k in sorted(set(numbers), reverse=True):
        entry = entries[k - 1]
        prs.part.drop_rel(entry.rId)
        ids.remove(entry)


def move(prs: Any, n: int, to: int) -> None:
    _check(prs, n, to)
    ids = _slide_ids(prs)
    entry = list(ids)[n - 1]
    ids.remove(entry)
    ids.insert(to - 1, entry)


def listing(path: Path) -> list[str]:
    import deck_model  # the QA scripts' title logic: skips section trackers and footers

    return [
        f"{s.index:>3}  [{s.layout or '-'}]  {s.title.splitlines()[0][:80] if s.title else ''}"
        for s in deck_model.load(str(path)).slides
    ]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--output", help="write here instead of rewriting the input")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("list")
    p = sub.add_parser("dup")
    p.add_argument("n", type=int)
    p.add_argument("--after", type=int)
    p = sub.add_parser("delete")
    p.add_argument("n", type=int, nargs="+")
    p = sub.add_parser("move")
    p.add_argument("n", type=int)
    p.add_argument("--to", type=int, required=True)
    args = ap.parse_args(argv)

    src = Path(args.deck)
    if src.suffix.lower() != ".pptx" or not src.is_file():
        sys.exit(f"deck_edit: {src} is not a .pptx file")
    prs = Presentation(str(src))
    if args.cmd == "list":
        print("\n".join(listing(src)))
        return 0
    if args.cmd == "dup":
        if args.after is not None:
            _check(prs, args.after)
        at = duplicate(prs, args.n, args.after)
        print(f"deck_edit: slide {args.n} copied to position {at}")
    elif args.cmd == "delete":
        delete(prs, args.n)
        print(f"deck_edit: deleted slide(s) {', '.join(map(str, sorted(set(args.n))))}")
    elif args.cmd == "move":
        move(prs, args.n, args.to)
        print(f"deck_edit: slide {args.n} moved to position {args.to}")
    out = Path(args.output) if args.output else src
    prs.save(str(out))
    print(f"deck_edit: wrote {out} ({len(prs.slides)} slides); run validate_pptx.py next")
    return 0


if __name__ == "__main__":
    sys.exit(main())
